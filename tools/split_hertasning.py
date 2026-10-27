import re, zipfile, xml.etree.ElementTree as ET
from pathlib import Path
src = Path(r'C:\Users\axlfo\AppData\Local\hermes\attachments\rincian_faktur_penjualan_dokteracmobil_260914210103 ver2.xlsx')
out = Path(r'D:\PROJECT\DRAC-APP\faktur-hertasning-raw.xlsx')
ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
def colnum(ref):
    m=re.match(r'([A-Z]+)',ref); n=0
    for c in m.group(1): n=n*26+ord(c)-64
    return n-1
def value(c):
    ts=[x.text or '' for x in c.findall('.//m:t',ns)]
    if ts: return ''.join(ts)
    v=c.find('m:v',ns); return v.text if v is not None else ''
with zipfile.ZipFile(src) as z:
    rows=[]
    for _,e in ET.iterparse(z.open('xl/worksheets/sheet1.xml'),events=('end',)):
        if e.tag.endswith('}row'):
            rows.append({colnum(c.attrib.get('r','A1')):value(c) for c in e.findall('m:c',ns)})
            e.clear()
header=next(r for r in rows if any(v.strip()=='Tanggal' for v in r.values()))
maxcol=max(header)
headers=[header.get(i,'') for i in range(maxcol+1)]
num_i=next(i for i,v in header.items() if v.strip()=='Nomor #')
data=[]
for r in rows:
    num=r.get(num_i,'').strip()
    if num and num!='Nomor #' and num.upper().startswith('H'):
        data.append([r.get(i,'') for i in range(maxcol+1)])
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
wb=Workbook(); ws=wb.active; ws.title='Hertasning'
for c,h in enumerate(headers,1):
    cell=ws.cell(1,c,h); cell.font=Font(bold=True); cell.fill=PatternFill('solid',fgColor='D9EAF7')
for rr,row in enumerate(data,2):
    for c,v in enumerate(row,1): ws.cell(rr,c,v)
ws.freeze_panes='A2'; ws.auto_filter.ref=f'A1:{ws.cell(len(data)+1,len(headers)).coordinate}'
for c in range(1,len(headers)+1): ws.column_dimensions[ws.cell(1,c).column_letter].width=16
wb.save(out)
print('output:',out)
print('detail_rows:',len(data))
print('unique_invoices:',len({row[num_i] for row in rows if row.get(num_i,'').strip().upper().startswith('H')}))
print('columns:',len(headers))
