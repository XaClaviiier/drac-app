import re, zipfile, xml.etree.ElementTree as ET
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
src=Path(r'C:\Users\axlfo\AppData\Local\hermes\attachments\rincian_faktur_penjualan_dokteracmobil_260914210103 ver2.xlsx')
ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
def colnum(ref):
 m=re.match(r'([A-Z]+)',ref); n=0
 for c in m.group(1): n=n*26+ord(c)-64
 return n-1
def value(c):
 t=[x.text or '' for x in c.findall('.//m:t',ns)]
 if t:return ''.join(t)
 v=c.find('m:v',ns); return v.text if v is not None else ''
with zipfile.ZipFile(src) as z:
 rows=[]
 for _,e in ET.iterparse(z.open('xl/worksheets/sheet1.xml'),events=('end',)):
  if e.tag.endswith('}row'):
   rows.append({colnum(c.attrib.get('r','A1')):value(c) for c in e.findall('m:c',ns)}); e.clear()
header=next(r for r in rows if any(v.strip()=='Tanggal' for v in r.values()))
maxcol=max(header); headers=[header.get(i,'') for i in range(maxcol+1)]
num_i=next(i for i,v in header.items() if v.strip()=='Nomor #')
for code,name in [('C','cakalang'),('D','perintis')]:
 data=[]
 for r in rows:
  num=r.get(num_i,'').strip()
  if num and num!='Nomor #' and num.upper().startswith(code): data.append([r.get(i,'') for i in range(maxcol+1)])
 out=Path(r'D:\PROJECT\DRAC-APP')/f'faktur-{name}-raw.xlsx'
 wb=Workbook(); ws=wb.active; ws.title=name.title()
 for c,h in enumerate(headers,1):
  x=ws.cell(1,c,h); x.font=Font(bold=True); x.fill=PatternFill('solid',fgColor='D9EAF7')
 for rr,row in enumerate(data,2):
  for c,v in enumerate(row,1): ws.cell(rr,c,v)
 ws.freeze_panes='A2'; ws.auto_filter.ref=f'A1:{ws.cell(len(data)+1,len(headers)).coordinate}'
 wb.save(out)
 print(name, 'detail_rows=',len(data), 'unique_invoices=',len({r[num_i] for r in rows if r.get(num_i,'').strip().upper().startswith(code)}), 'output=',out)
